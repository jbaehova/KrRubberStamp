"""Authored placements and prose are inputs, never renderer-generated decisions."""

import json
import zipfile
from xml.etree import ElementTree as ET

import pymupdf
import pytest
from openpyxl import load_workbook
from PIL import Image

from render.authored import render_authored
from render.documents import HP, restore_scenario


def plan(filename, fields=None, *, note=None, title=None, duplicate_of=None):
    item = {
        "filename": filename,
        "format": filename.rsplit(".", 1)[1],
        "title": title or "마감 직전 도착한 발주 확인서",
        "note": note
        or "현장 담당자가 추가 납품분을 확인했습니다. 먼저 받은 견적과 구분해서 검토하세요.",
    }
    if fields is not None:
        item["fields"] = fields
    if duplicate_of is not None:
        item["duplicate_of"] = duplicate_of
    return item


def mapping(directory):
    return json.loads((directory / "extraction_map.json").read_text())


def test_explicit_document_order_and_nested_placements_roundtrip(tmp_path):
    scenario = {
        "employee": {"name": "가상 김서연", "salary": 3125000},
        "people": [{"age": 7, "eligible": True}, {"age": 22, "eligible": False}],
        "documents": [
            {"items": [{"name": "보관함", "quantity": 3}, {"name": "색인표", "quantity": 12}]}
        ],
        "empty": [],
        "missing": None,
    }
    documents = [
        plan("03_확인.hwpx", ["people.1", "documents.0.items.1", "empty"]),
        plan("01_계약.pdf", ["employee", "people.0"]),
        plan("04_전달.xlsx", ["documents.0.items.0", "missing", "employee.name"]),
        plan(
            "02_사본.png",
            duplicate_of="01_계약.pdf",
            note="팩스로 다시 받은 계약서입니다. 급여를 두 번 더하지 마세요.",
        ),
    ]
    inputs = render_authored(scenario, tmp_path, "D_extract", documents, seed=721)
    assert [item["path"] for item in inputs] == ["inputs/" + item["filename"] for item in documents]
    assert restore_scenario(tmp_path) == scenario
    source_map = mapping(tmp_path)
    assert {
        item["file"] for item in source_map["records"] if item["path"][:2] == ["people", 1]
    } == {"inputs/03_확인.hwpx"}
    first_records = [
        item["path"] for item in source_map["records"] if item["file"] == "inputs/03_확인.hwpx"
    ]
    assert first_records == [
        ["people", 1, "age"],
        ["people", 1, "eligible"],
        ["documents", 0, "items", 1, "name"],
        ["documents", 0, "items", 1, "quantity"],
        ["empty"],
    ]
    assert [item["path"] for item in source_map["records"]].count(["employee", "name"]) == 2
    assert all(set(item) == {"path", "kind", "file", "selector"} for item in source_map["records"])
    assert all("value" not in item for item in source_map["records"])
    assert set(source_map) == {"version", "containers", "records", "documents"}
    with Image.open(tmp_path / "inputs/02_사본.png") as image:
        assert image.getextrema()[0] < 100
        assert image.height > 1600


@pytest.mark.parametrize(
    "bad_fields,message",
    [
        (["salary"], "uncovered"),
        (["salary", "people.3"], "Unknown"),
        (["salary", "people.00"], "Unknown"),
        ([], "no selected fields"),
    ],
)
def test_missing_or_unknown_selection_is_rejected_before_output(tmp_path, bad_fields, message):
    with pytest.raises(ValueError, match=message):
        render_authored(
            {"salary": 100, "people": [{"age": 9}]},
            tmp_path,
            "B_payroll",
            [plan("01.pdf", bad_fields)],
        )
    assert not (tmp_path / "inputs").exists()
    assert not (tmp_path / "extraction_map.json").exists()


def test_all_primary_formats_preserve_complete_authored_narrative(tmp_path):
    title = "도서관 별관 개관 준비를 위해 다시 확인한 물품별 공급 조건과 정산 기록"
    note = (
        "담당자는 별관에 먼저 도착한 물품을 본관 발주분과 분리해 두었습니다. " * 18
        + "최종 전달 사항은 분할 납품의 잔여 수량 확인입니다."
    )
    scenario = {"first": 210000, "second": 13, "third": "가상 별관"}
    documents = [
        plan("01.pdf", ["first"], note=note, title=title),
        plan("02.xlsx", ["second"], note=note, title=title),
        plan("03.hwpx", ["third"], note=note, title=title),
    ]
    render_authored(scenario, tmp_path, "D_extract", documents)
    assert restore_scenario(tmp_path) == scenario
    compact_note = "".join(note.split())
    with pymupdf.open(tmp_path / "inputs/01.pdf") as pdf:
        text = "".join(page.get_text() for page in pdf)
        assert compact_note in "".join(text.split())
        assert "".join(title.split()) in "".join(text.split())
        for page in pdf:
            for block in page.get_text("dict")["blocks"]:
                if "lines" in block:
                    assert block["bbox"][2] < page.rect.width
                    assert block["bbox"][3] < page.rect.height
    workbook = load_workbook(tmp_path / "inputs/02.xlsx")
    sheet = workbook["원자료"]
    assert sheet["A1"].value == title
    note_cells = [
        sheet.cell(index, 1)
        for index in range(3, sheet.max_row)
        if isinstance(sheet.cell(index, 1).value, str)
    ]
    assert compact_note in "".join("".join(cell.value.split()) for cell in note_cells)
    assert sheet.row_dimensions[3].height > 30
    assert all(
        sheet.row_dimensions[cell.row].height <= 409.5 for cell in note_cells if cell.row != 5
    )
    workbook.close()
    with zipfile.ZipFile(tmp_path / "inputs/03.hwpx") as archive:
        section = ET.fromstring(archive.read("Contents/section0.xml"))
        text = "".join(node.text or "" for node in section.iter(f"{{{HP}}}t"))
        assert compact_note in "".join(text.split())


def test_large_workbook_narrative_keeps_fields_at_shifted_coordinates(tmp_path):
    note = (
        "교체 대상으로 표시한 장비는 재고 대장에 남아 있지만 사용할 수 없는 상태입니다. " * 100
    ) + "끝까지 기재한 별도 확인 사항입니다."
    render_authored(
        {"quantity": 19}, tmp_path, "D_extract", [plan("01.xlsx", ["quantity"], note=note)]
    )
    assert restore_scenario(tmp_path) == {"quantity": 19}
    record = mapping(tmp_path)["records"][0]
    assert int(record["selector"]["cell"][1:]) > 6
    workbook = load_workbook(tmp_path / "inputs/01.xlsx")
    sheet = workbook["원자료"]
    assert all(
        dimension.height is None or dimension.height <= 409.5
        for dimension in sheet.row_dimensions.values()
    )
    text = "".join(str(sheet.cell(index, 1).value or "") for index in range(3, sheet.max_row))
    assert "끝까지기재한별도확인사항입니다." in "".join(text.split())
    workbook.close()


def test_visible_source_changes_the_restored_fact(tmp_path):
    render_authored({"amount": 12345}, tmp_path, "D_extract", [plan("01.xlsx", ["amount"])])
    record = mapping(tmp_path)["records"][0]
    source = tmp_path / record["file"]
    workbook = load_workbook(source)
    workbook["원자료"][record["selector"]["cell"]] = 22222
    workbook.save(source)
    workbook.close()
    assert restore_scenario(tmp_path) == {"amount": 22222}


def test_seed_does_not_change_facts_placements_titles_or_notes(tmp_path, monkeypatch):
    import scenarios.b_payroll

    def forbidden(*args, **kwargs):
        pytest.fail("An authored renderer must never generate its scenario")

    monkeypatch.setattr(scenarios.b_payroll, "generate", forbidden)
    scenario = {"pay": 90123, "employee": "가상 이소미"}
    documents = [plan("01.pdf", ["employee"]), plan("02.xlsx", ["pay"])]
    first, second = tmp_path / "first", tmp_path / "second"
    render_authored(scenario, first, "B_payroll", documents, seed=1)
    render_authored(scenario, second, "B_payroll", documents, seed=99999)
    assert mapping(first) == mapping(second)
    assert restore_scenario(first) == restore_scenario(second) == scenario
    assert (first / "inputs/01.pdf").read_bytes() == (second / "inputs/01.pdf").read_bytes()


@pytest.mark.parametrize(
    "documents,message",
    [
        ([plan("../escape.pdf", ["amount"])], "Unsafe"),
        ([plan("01.pdf", ["amount"]), plan("01.pdf", ["amount"])], "Duplicate"),
        ([plan("01.png", duplicate_of="02.pdf"), plan("02.pdf", ["amount"])], "earlier PDF"),
        (
            [plan("01.pdf", ["amount"]), plan("02.png", ["amount"], duplicate_of="01.pdf")],
            "cannot be the source",
        ),
    ],
)
def test_invalid_file_and_duplicate_plans_are_rejected(tmp_path, documents, message):
    with pytest.raises(ValueError, match=message):
        render_authored({"amount": 12}, tmp_path, "D_extract", documents)
    assert not (tmp_path / "inputs").exists()


def test_existing_inputs_are_preserved_and_not_overwritten(tmp_path):
    (tmp_path / "inputs").mkdir()
    original = tmp_path / "inputs/01.pdf"
    original.write_bytes(b"existing private source")
    with pytest.raises(ValueError, match="empty"):
        render_authored({"amount": 12}, tmp_path, "D_extract", [plan("01.pdf", ["amount"])])
    assert original.read_bytes() == b"existing private source"


def test_multiline_facts_require_the_author_to_choose_lossless_format(tmp_path):
    scenario = {"message": "첫 번째 납품\n두 번째 납품\t별도 확인"}
    with pytest.raises(ValueError, match="require XLSX or HWPX"):
        render_authored(scenario, tmp_path, "D_extract", [plan("01.pdf", ["message"])])
    render_authored(scenario, tmp_path, "D_extract", [plan("01.hwpx", ["message"])])
    assert restore_scenario(tmp_path) == scenario
