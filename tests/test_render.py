"""Source document tests exercise actual data cells, not stored scenarios."""

import json
import zipfile
from importlib import import_module
from xml.etree import ElementTree as ET

import pymupdf
import pytest
from openpyxl import load_workbook
from PIL import Image

from render import render_task, restore_scenario
from render.documents import HP, HH, OPF, TEMPLATES


@pytest.mark.parametrize("domain", ["A_yearend", "B_payroll", "C_vat", "D_extract"])
@pytest.mark.parametrize("difficulty,count", [("easy", 2), ("medium", 5), ("hard", 10)])
def test_actual_domain_input_roundtrip(tmp_path, domain, difficulty, count):
    scenario = import_module("scenarios." + domain.lower()).generate(19293, difficulty)
    files = render_task(scenario, tmp_path, domain, difficulty, 19293)
    assert len(files) == count
    assert restore_scenario(tmp_path) == scenario
    assert set((tmp_path / "inputs").iterdir()) == {tmp_path / item["path"] for item in files}
    for item in files:
        path = tmp_path / item["path"]
        assert path.stat().st_size > 100
        if item["format"] == "pdf":
            with pymupdf.open(path) as pdf:
                assert sum(len(page.get_text()) for page in pdf) > 100
        if item["format"] == "png":
            assert item["duplicate_of"]
            with Image.open(path) as image:
                assert image.getextrema()[0] < 200
                assert image.width > 900
    mapping = json.loads((tmp_path / "extraction_map.json").read_text())
    for record in mapping["records"]:
        assert set(record) == {"path", "kind", "file", "selector"}
    assert "scenario" not in mapping and "gold" not in mapping and "trace" not in mapping


def test_numeric_cells_are_numeric_and_string_cells_are_not_formulas(tmp_path):
    scenario = {
        "n0": 17,
        "value": 4200000,
        "formula": "=SUM(1,2)",
        "empty": "",
        "unicode": "가상㈜ 거래처 ①",
        "multiline": "첫 줄\n둘째 줄\t끝",
    }
    render_task(scenario, tmp_path, "B_payroll", "easy", 11)
    mapping = json.loads((tmp_path / "extraction_map.json").read_text())
    numeric_records = [
        r for r in mapping["records"] if r["file"].endswith(".xlsx") and r["kind"] == "int"
    ]
    assert numeric_records
    wb = load_workbook(tmp_path / numeric_records[0]["file"])
    r = numeric_records[0]
    assert isinstance(wb[r["selector"]["sheet"]][r["selector"]["cell"]].value, int)
    wb.close()
    assert restore_scenario(tmp_path) == scenario


def test_changing_visible_source_cell_changes_restored_value(tmp_path):
    scenario = {"amount": 120000, "quantity": 12, "name": "가상물품"}
    render_task(scenario, tmp_path, "B_payroll", "easy", 3)
    mapping = json.loads((tmp_path / "extraction_map.json").read_text())
    record = next(
        r for r in mapping["records"] if r["file"].endswith(".xlsx") and r["kind"] == "int"
    )
    workbook_path = tmp_path / record["file"]
    wb = load_workbook(workbook_path)
    wb[record["selector"]["sheet"]][record["selector"]["cell"]] = 999
    wb.save(workbook_path)
    wb.close()
    restored = restore_scenario(tmp_path)
    assert restored != scenario
    assert restored[record["path"][0]] == 999


def test_missing_source_and_invalid_type_are_detected(tmp_path):
    render_task({"a": 1, "flag": True}, tmp_path, "B_payroll", "easy", 3)
    mapping = json.loads((tmp_path / "extraction_map.json").read_text())
    record = next(r for r in mapping["records"] if r["kind"] == "bool")
    path = tmp_path / record["file"]
    wb = load_workbook(path)
    wb[record["selector"]["sheet"]][record["selector"]["cell"]] = "garbled"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="boolean"):
        restore_scenario(tmp_path)
    path.unlink()
    with pytest.raises(FileNotFoundError):
        restore_scenario(tmp_path)


def test_hwpx_package_table_and_style_references(tmp_path):
    scenario = {
        "amount": 1250000,
        "name": "가상종이업체",
        "quantity": 15,
        "items": [{"name": "복사용지", "unit_price": 9000}],
        "vat": 125000,
        "note": "합성자료",
    }
    files = render_task(scenario, tmp_path, "D_extract", "medium", 21)
    path = tmp_path / next(item["path"] for item in files if item["format"] == "hwpx")
    with zipfile.ZipFile(path) as archive:
        assert archive.infolist()[0].filename == "mimetype"
        assert archive.infolist()[0].compress_type == zipfile.ZIP_STORED
        assert archive.read("mimetype") == b"application/hwp+zip"
        assert archive.testzip() is None
        for member in archive.namelist():
            if member.endswith((".xml", ".hpf")):
                ET.fromstring(archive.read(member))
        head = ET.fromstring(archive.read("Contents/header.xml"))
        section = ET.fromstring(archive.read("Contents/section0.xml"))
        manifest = ET.fromstring(archive.read("Contents/content.hpf"))
        assert [
            item.get("idref") for item in manifest.findall(f"{{{OPF}}}spine/{{{OPF}}}itemref")
        ] == ["header", "section0"]
        assert head.find(f"{{{HH}}}refList/{{{HH}}}charProperties/{{{HH}}}charPr").get("id") == "0"
        tables = section.findall(f".//{{{HP}}}tbl")
        assert len(tables) == 1
        table = tables[0]
        rows = table.findall(f"{{{HP}}}tr")
        assert int(table.get("rowCnt")) == len(rows) > 1
        assert all(len(row.findall(f"{{{HP}}}tc")) == 2 for row in rows)
        assert section.find(f".//{{{HP}}}secPr/{{{HP}}}pagePr") is not None
    assert restore_scenario(tmp_path) == scenario


def test_pdf_contains_no_embedded_scenario_or_answer(tmp_path):
    render_task({"name": "가상 근로자", "salary": 4800000}, tmp_path, "B_payroll", "easy", 1)
    with pymupdf.open(tmp_path / "inputs/01_source.pdf") as pdf:
        assert pdf.embfile_count() == 0
        assert pdf.metadata["subject"] == ""
        text = pdf[0].get_text()
        assert "가상 근로자" in text
        assert "4,800,000" not in pdf.metadata.values()
        assert "KrRubberStamp" in text
    assert not (tmp_path / "inputs/extraction_map.json").exists()


def test_five_templates_and_transactions_stay_together(tmp_path):
    assert len(TEMPLATES) == 5
    scenario = {
        "documents": [
            {
                "document_id": f"SYN-{i}",
                "vendor": f"가상업체{i}",
                "items": [{"name": "복사용지", "quantity": i + 1, "unit_price": 1000}],
            }
            for i in range(8)
        ],
        "exceptions": ["중복 증빙 확인", "단가 부가세 확인"],
    }
    render_task(scenario, tmp_path, "D_extract", "hard", 0)
    mapping = json.loads((tmp_path / "extraction_map.json").read_text())
    assert {d["template_variant"] for d in mapping["documents"]} == set(range(5))
    for index in range(8):
        assert (
            len({r["file"] for r in mapping["records"] if r["path"][:2] == ["documents", index]})
            == 1
        )
    assert restore_scenario(tmp_path) == scenario


def test_same_seed_pdf_and_scan_are_reproducible(tmp_path):
    first, second = tmp_path / "one", tmp_path / "two"
    scenario = {"name": "가상업체", "amount": 120000, "items": [{"quantity": 12}]}
    one = render_task(scenario, first, "D_extract", "easy", 37)
    two = render_task(scenario, second, "D_extract", "easy", 37)
    assert one == two
    for item in one:
        assert (first / item["path"]).read_bytes() == (second / item["path"]).read_bytes()
