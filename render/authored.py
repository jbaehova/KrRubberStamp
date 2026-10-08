"""Render an author's explicit source documents without assigning any facts."""

from __future__ import annotations

import json
import os
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from render.documents import (
    FONT_NAME,
    FONT_PATH,
    HP,
    PAGE_HEIGHT,
    PAGE_WIDTH,
    TEMPLATES,
    _display,
    _flatten,
    _font,
    _hwpx,
    _label,
    _scan,
    _wrapped,
    _xlsx,
)


def _prefix(scenario: dict, selector: str) -> tuple[str | int, ...]:
    """Resolve list indexes against the actual tree, rather than guessing keys."""
    if not isinstance(selector, str) or not selector or any(not x for x in selector.split(".")):
        raise ValueError(f"Invalid authored field selector: {selector!r}")
    value: Any = scenario
    path: list[str | int] = []
    for token in selector.split("."):
        if isinstance(value, dict) and token in value:
            key: str | int = token
        elif isinstance(value, list) and token.isascii() and token.isdecimal():
            key = int(token)
            if str(key) != token or key >= len(value):
                raise ValueError(f"Unknown authored field selector: {selector}")
        else:
            raise ValueError(f"Unknown authored field selector: {selector}")
        path.append(key)
        value = value[key]
    return tuple(path)


def _safe_filename(filename: Any, fmt: str) -> str:
    if (
        not isinstance(filename, str)
        or not filename
        or filename in (".", "..")
        or any(char in filename for char in ("/", "\\", "\x00", "\n", "\r"))
        or Path(filename).name != filename
        or Path(filename).suffix != "." + fmt
    ):
        raise ValueError(f"Unsafe authored filename or extension mismatch: {filename!r}")
    return filename


def _prepare(scenario: dict, documents: list[dict], domain: str) -> tuple[list, list]:
    if not isinstance(scenario, dict) or not scenario:
        raise ValueError("Authored scenario must be a nonempty object")
    if not isinstance(documents, list) or not documents:
        raise ValueError("An authored document plan is required")
    containers: list[dict] = []
    records: list[dict] = []
    _flatten(scenario, [], containers, records)
    for record in records:
        record["label"] = _label(record["path"], domain)
    prepared: list[dict] = []
    covered: set[tuple] = set()
    formats: dict[str, str] = {}
    allowed = {"filename", "format", "title", "note", "fields", "template_variant", "duplicate_of"}
    for document in documents:
        if not isinstance(document, dict) or set(document) - allowed:
            raise ValueError("Authored document plan has unknown properties")
        fmt = document.get("format")
        if fmt not in ("pdf", "xlsx", "hwpx", "png"):
            raise ValueError(f"Unsupported authored source format: {fmt!r}")
        filename = _safe_filename(document.get("filename"), fmt)
        if filename in formats:
            raise ValueError(f"Duplicate authored filename: {filename}")
        title, note = document.get("title"), document.get("note")
        if not isinstance(title, str) or not title.strip() or len(title) > 60:
            raise ValueError("Each authored title must contain 1 to 60 characters")
        if not isinstance(note, str) or not note.strip():
            raise ValueError("Each authored document needs its own narrative note")
        if any(char in title + note for char in ("\x00", "\r", "\t")):
            raise ValueError("Authored title and note contain unsupported control characters")
        variant = document.get("template_variant", 0)
        if type(variant) is not int or not 0 <= variant < len(TEMPLATES):
            raise ValueError("Authored template_variant must be an integer from 0 to 4")
        entry = {**document, "template_variant": variant, "records": []}
        if fmt == "png":
            duplicate = document.get("duplicate_of")
            if not isinstance(duplicate, str) or formats.get(duplicate) != "pdf":
                raise ValueError("Authored PNG duplicate_of must name an earlier PDF")
            if document.get("fields"):
                raise ValueError("Authored PNG copies cannot be the source of scenario fields")
        else:
            if "duplicate_of" in document:
                raise ValueError("Only an authored PNG can declare duplicate_of")
            fields = document.get("fields")
            if not isinstance(fields, list) or not fields:
                raise ValueError(f"Authored source {filename} has no selected fields")
            for selector in fields:
                prefix = _prefix(scenario, selector)
                matches = [r for r in records if tuple(r["path"][: len(prefix)]) == prefix]
                if not matches:
                    raise ValueError(f"Unknown authored field selector: {selector}")
                entry["records"].extend(matches)
                covered.update(tuple(record["path"]) for record in matches)
            if fmt == "pdf" and any(
                r["kind"] == "str" and any(char in r["value"] for char in ("\n", "\r", "\t"))
                for r in entry["records"]
            ):
                raise ValueError(f"Multiline authored facts in {filename} require XLSX or HWPX")
        formats[filename] = fmt
        prepared.append(entry)
    missing = [".".join(map(str, r["path"])) for r in records if tuple(r["path"]) not in covered]
    if missing:
        raise ValueError(
            "Authored document plan leaves source fields uncovered: " + ", ".join(missing)
        )
    return containers, prepared


def _authored_pdf(
    path: Path, records: list[dict], title: str, variant: int, note: str
) -> list[dict]:
    """Wrap the complete narrative before the table, with measured PDF selectors."""
    _font()
    design = TEMPLATES[variant]
    accent = colors.HexColor("#" + design["accent"])
    pdf = canvas.Canvas(
        str(path), pagesize=(PAGE_WIDTH, PAGE_HEIGHT), invariant=1, pageCompression=1
    )
    pdf.setTitle(title)
    pdf.setAuthor("KrRubberStamp authored office case")
    pdf.setSubject("")
    margin, label_width, size = 38, 272, 9.4
    x_value = margin + label_width
    width = PAGE_WIDTH - 2 * margin
    value_width = PAGE_WIDTH - margin - x_value
    page_number = 0
    selectors: list[dict] = []

    def heading() -> float:
        y = PAGE_HEIGHT - 47
        pdf.setFillColor(accent)
        pdf.setFont(FONT_NAME, 17)
        for line in _wrapped(title, width, 17):
            pdf.drawString(margin, y, line)
            y -= 23
        pdf.setFillColor(colors.HexColor("#525B66"))
        pdf.setFont(FONT_NAME, 8)
        pdf.drawString(margin, y - 2, "사내 업무용 자료 | 모든 이름과 식별번호는 합성 정보입니다")
        pdf.setFont(FONT_NAME, 7)
        pdf.drawString(margin, 22, f"KrRubberStamp | {page_number + 1}쪽")
        return y - 25

    def new_page() -> float:
        nonlocal page_number
        pdf.showPage()
        page_number += 1
        return heading()

    def table_header(y: float) -> float:
        pdf.setFillColor(colors.HexColor("#" + design["header"]))
        pdf.rect(margin, y - 24, width, 24, fill=1, stroke=0)
        pdf.setFillColor(accent)
        pdf.setFont(FONT_NAME, 9)
        pdf.drawString(margin + 7, y - 17, "자료 항목")
        pdf.drawString(x_value + 7, y - 17, "원자료 기재 내용")
        return y - 24

    y = heading()
    for line in _wrapped(note, width, 9.5):
        if y < 62:
            y = new_page()
        pdf.setFillColor(colors.HexColor("#202A34"))
        pdf.setFont(FONT_NAME, 9.5)
        pdf.drawString(margin, y, line)
        y -= 15
    y -= 12
    if y < 105:
        y = new_page()
    y = table_header(y)
    for record in records:
        labels = _wrapped(record["label"], label_width - 14, size)
        values = _wrapped(_display(record), value_width - 14, size)
        row_height = max(27, 14 + 13 * max(len(labels), len(values)))
        if row_height > PAGE_HEIGHT - 180:
            raise ValueError("One authored field is too large to fit on a PDF page")
        if y - row_height < 43:
            y = table_header(new_page())
        if y - row_height < 43:
            raise ValueError("One authored field is too large to fit below its PDF title")
        if variant in (0, 1, 4) and len(selectors) % 2:
            pdf.setFillColor(colors.HexColor("#F8F9FA"))
            pdf.rect(margin, y - row_height, width, row_height, stroke=0, fill=1)
        pdf.setStrokeColor(colors.HexColor("#CCD2D8"))
        pdf.setLineWidth(0.4)
        pdf.line(margin, y - row_height, PAGE_WIDTH - margin, y - row_height)
        if variant in (0, 2, 3):
            pdf.line(x_value, y, x_value, y - row_height)
        pdf.setFillColor(colors.HexColor("#202A34"))
        pdf.setFont(FONT_NAME, size)
        for index, line in enumerate(labels):
            pdf.drawString(margin + 7, y - 16 - index * 13, line)
        for index, line in enumerate(values):
            pdf.drawString(x_value + 7, y - 16 - index * 13, line)
        selectors.append(
            {
                "page": page_number,
                "bbox": [
                    round(x_value + 3, 3),
                    round(PAGE_HEIGHT - y, 3),
                    round(PAGE_WIDTH - margin - 3, 3),
                    round(PAGE_HEIGHT - y + row_height, 3),
                ],
            }
        )
        y -= row_height
    pdf.save()
    return selectors


def _rewrite_zip(path: Path, replacements: dict[str, bytes] | None = None) -> None:
    replacements = replacements or {}
    with zipfile.ZipFile(path) as archive:
        members = [
            (item, replacements.get(item.filename, archive.read(item.filename)))
            for item in archive.infolist()
        ]
    with zipfile.ZipFile(path, "w") as archive:
        for item, content in members:
            item.date_time = (2026, 1, 1, 0, 0, 0)
            archive.writestr(item, content)


def _authored_xlsx(
    path: Path, records: list[dict], title: str, variant: int, note: str
) -> list[dict]:
    _font()
    selectors = _xlsx(path, records, title, variant, "")
    workbook = load_workbook(path)
    try:
        sheet = workbook["원자료"]
        title_lines = _wrapped(title, 600, 18)
        note_lines = _wrapped(note, 600, 10)
        chunks = [note_lines[index : index + 22] for index in range(0, len(note_lines), 22)]
        added = len(chunks) - 1
        if added:
            sheet.insert_rows(4, amount=added)
        sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        sheet["A1"].data_type = "s"
        sheet.row_dimensions[1].height = max(34, 24 * len(title_lines))
        for index, lines in enumerate(chunks, start=3):
            if index > 3:
                sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=3)
            sheet.cell(index, 1, note if len(chunks) == 1 else "\n".join(lines))
            sheet.cell(index, 1).data_type = "s"
            sheet.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.cell(index, 1).font = Font(name="Noto Sans KR", size=10, color="202A34")
            sheet.row_dimensions[index].height = max(30, 17 * len(lines))
        for selector, record in zip(selectors, records, strict=True):
            old_row = int(selector["cell"][1:])
            selector["cell"] = "C" + str(old_row + added)
            label_height = len(_wrapped(record["label"], 330, 10))
            value = sheet[selector["cell"]].value
            value_height = len(_wrapped(str(value), 205, 10))
            height = max(26, 17 * max(label_height, value_height))
            if height > 409.5:
                raise ValueError("One authored field is too large to fit in an XLSX row")
            sheet.row_dimensions[old_row + added].height = height
        sheet.freeze_panes = "C" + str(6 + added)
        sheet.auto_filter.ref = f"A{5 + added}:C{5 + added + len(records)}"
        sheet.print_title_rows = f"{5 + added}:{5 + added}"
        workbook.save(path)
    finally:
        workbook.close()
    _rewrite_zip(path)
    return selectors


def _authored_hwpx(
    path: Path, records: list[dict], title: str, variant: int, note: str
) -> list[dict]:
    selectors = _hwpx(path, records, title, variant, note)
    with zipfile.ZipFile(path) as archive:
        section = ET.fromstring(archive.read("Contents/section0.xml"))
    paragraphs = section.findall(f"{{{HP}}}p")
    original = paragraphs[2]
    font_size = (10, 10.5, 9.5, 11, 10)[variant]
    lines = _wrapped(note, 475, font_size)
    next_id = 1 + max(int(p.get("id", "0")) for p in section.iter(f"{{{HP}}}p"))
    section.remove(original)
    for index, line in enumerate(lines):
        paragraph = deepcopy(original)
        paragraph.set("id", str(next_id + index))
        paragraph.find(f".//{{{HP}}}t").text = line
        section.insert(2 + index, paragraph)
    _rewrite_zip(
        path,
        {"Contents/section0.xml": ET.tostring(section, encoding="utf-8", xml_declaration=True)},
    )
    return selectors


def _authored_scan(source: Path, path: Path, title: str, note: str, seed: int) -> None:
    _scan(source, path, seed)
    with Image.open(path) as original:
        original.load()
        title_font = ImageFont.truetype(str(FONT_PATH), 32)
        note_font = ImageFont.truetype(str(FONT_PATH), 20)
        probe = ImageDraw.Draw(original)

        def wrap(text: str, font: ImageFont.FreeTypeFont) -> list[str]:
            lines: list[str] = []
            for paragraph in text.split("\n"):
                current = ""
                for char in paragraph:
                    if (
                        current
                        and probe.textlength(current + char, font=font) > original.width - 120
                    ):
                        lines.append(current)
                        current = ""
                    current += char
                lines.append(current)
            return lines

        title_lines, note_lines = wrap(title, title_font), wrap(note, note_font)
        header_height = 55 + 44 * len(title_lines) + 30 * len(note_lines)
        image = Image.new("L", (original.width, original.height + header_height), 255)
        image.paste(original, (0, header_height))
        draw = ImageDraw.Draw(image)
        y = 24
        for line in title_lines:
            draw.text((60, y), line, font=title_font, fill=35)
            y += 44
        for line in note_lines:
            draw.text((60, y), line, font=note_font, fill=50)
            y += 30
        image.save(path, optimize=True)


def render_authored(
    scenario: dict, task_dir: Path, domain: str, documents: list[dict], seed: int = 0
) -> list[dict]:
    """Write exactly the source documents and placements supplied by the author.

    ``seed`` affects only the requested duplicate scans. No scenario facts,
    document count, ordering, titles, notes or field assignments are generated.
    """
    containers, prepared = _prepare(scenario, documents, domain)
    task_dir = Path(task_dir)
    input_dir = task_dir / "inputs"
    map_path = task_dir / "extraction_map.json"
    if (
        task_dir.is_symlink()
        or input_dir.is_symlink()
        or map_path.exists()
        or map_path.is_symlink()
    ):
        raise ValueError("Authored output would overwrite a map or follow a symlink")
    if input_dir.exists() and (not input_dir.is_dir() or any(input_dir.iterdir())):
        raise ValueError("Authored inputs directory must be empty")
    task_dir.mkdir(parents=True, exist_ok=True)
    inputs: list[dict] = []
    mapped: list[dict] = []
    with tempfile.TemporaryDirectory(prefix=".authored-", dir=task_dir) as temporary:
        staged = Path(temporary)
        for index, document in enumerate(prepared):
            filename, fmt = document["filename"], document["format"]
            path = staged / filename
            entry = {"path": "inputs/" + filename, "format": fmt}
            if fmt == "png":
                _authored_scan(
                    staged / document["duplicate_of"],
                    path,
                    document["title"],
                    document["note"],
                    seed * 31 + index,
                )
                entry["duplicate_of"] = "inputs/" + document["duplicate_of"]
            else:
                writer = {"pdf": _authored_pdf, "xlsx": _authored_xlsx, "hwpx": _authored_hwpx}[fmt]
                selectors = writer(
                    path,
                    document["records"],
                    document["title"],
                    document["template_variant"],
                    document["note"],
                )
                for record, selector in zip(document["records"], selectors, strict=True):
                    mapped.append(
                        {
                            "path": record["path"],
                            "kind": record["kind"],
                            "file": entry["path"],
                            "selector": selector,
                        }
                    )
            inputs.append(entry)
        mapping = {
            "version": 1,
            "containers": containers,
            "records": mapped,
            "documents": [
                {**entry, "template_variant": document["template_variant"]}
                for entry, document in zip(inputs, prepared, strict=True)
            ],
        }
        staged_map = staged / "extraction_map.json"
        staged_map.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
        input_dir.mkdir(exist_ok=True)
        # Hard linking on the same filesystem refuses an existing destination.
        for entry in inputs:
            os.link(staged / Path(entry["path"]).name, task_dir / entry["path"])
        os.link(staged_map, map_path)
    return inputs
