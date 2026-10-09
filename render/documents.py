"""Deterministic Korean PDF, XLSX, HWPX and scan source documents."""

from __future__ import annotations

import json
import math
import random
import zipfile
from datetime import datetime
from importlib import import_module
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import pymupdf as fitz
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_PATH = Path(__file__).resolve().parents[1] / "assets/fonts/NotoSansKR-Regular.ttf"
FONT_NAME = "KrRubberStamp-NotoSansKR"
PAGE_WIDTH, PAGE_HEIGHT = A4
EMPTY_LABEL = {"null": "해당 없음", "list": "항목 없음", "dict": "기재 항목 없음"}
TEMPLATES = (
    {"name": "청색 표준", "accent": "244B75", "header": "EEF3F8", "style": 0},
    {"name": "녹색 장부", "accent": "28594E", "header": "EDF5F0", "style": 1},
    {"name": "흑백 공문", "accent": "30353B", "header": "F0F0F0", "style": 2},
    {"name": "갈색 거래표", "accent": "755339", "header": "F7F1E9", "style": 3},
    {"name": "자주색 실무", "accent": "654268", "header": "F3EDF4", "style": 4},
)
DOC_TITLES = {
    "A_yearend": (
        "근로소득 자료 확인서",
        "공제 증빙 제출내역",
        "연말정산 참고자료",
        "부양가족 및 지출 확인표",
        "납세 자료 대조표",
    ),
    "B_payroll": (
        "근로계약 및 지급조건",
        "근태 및 수당 정산자료",
        "급여 규정 확인서",
        "보험료 적용 기초자료",
        "월 지급자료 대조표",
    ),
    "C_vat": (
        "매출 거래 확인서",
        "매입 증빙 정리표",
        "거래 조건 확인서",
        "신고 기초자료 명세",
        "세무 자료 대조표",
    ),
    "D_extract": (
        "거래명세서",
        "품목별 견적자료",
        "거래 조건 확인서",
        "업체별 납품내역",
        "증빙 대조표",
    ),
}
FALLBACK_LABELS = {
    "domain": "업무 구분",
    "year": "귀속 연도",
    "month": "지급 월",
    "employee": "근로자 정보",
    "name": "이름",
    "company": "업체",
    "company_name": "업체명",
    "vendor": "거래처",
    "vendor_name": "거래처명",
    "business_number": "사업자등록번호",
    "employee_name": "근로자명",
    "resident_number": "주민등록번호",
    "salary": "급여",
    "annual_salary": "연간 급여",
    "gross_salary": "총급여",
    "base_salary": "기본급",
    "base_pay": "기본급",
    "dependents": "부양가족",
    "children": "자녀",
    "age": "나이",
    "income": "소득",
    "sales": "매출",
    "purchases": "매입",
    "items": "품목",
    "transactions": "거래",
    "quantity": "수량",
    "unit_price": "단가",
    "amount": "금액",
    "vat": "부가세",
    "supply_amount": "공급가액",
    "supply_price": "공급가액",
    "tax": "세액",
    "date": "일자",
    "description": "내용",
    "category": "구분",
    "hours": "시간",
    "overtime_hours": "연장근로 시간",
    "night_hours": "야간근로 시간",
    "holiday_hours": "휴일근로 시간",
    "exceptions": "확인 필요사항",
    "exception_type": "예외 구분",
    "source_id": "증빙 식별번호",
    "document_id": "문서 번호",
    "document_status_records": "원본 효력 확인 기록",
    "processing_date": "원본 대조 마감일",
    "state": "확정 효력 상태",
    "confirmed_on": "효력 확인일",
    "replaces_document_id": "대체한 취소 원본 번호",
    "taxable": "과세 여부",
    "deductible": "공제 가능 여부",
    "duplicate_of": "중복 원본 번호",
    "notes": "비고",
    "note": "비고",
    "id": "번호",
    "type": "종류",
    "paid_tax": "기납부세액",
    "withheld_tax": "원천징수 소득세",
    "insurance": "보험료",
    "medical": "의료비",
    "education": "교육비",
    "donations": "기부금",
    "rent": "월세",
    "cards": "카드 사용내역",
}
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HS = "http://www.hancom.co.kr/hwpml/2011/section"
HH = "http://www.hancom.co.kr/hwpml/2011/head"
HC = "http://www.hancom.co.kr/hwpml/2011/core"
OPF = "http://www.idpf.org/2007/opf/"
for prefix, uri in (("hp", HP), ("hs", HS), ("hh", HH), ("hc", HC), ("opf", OPF)):
    ET.register_namespace(prefix, uri)


def _font() -> None:
    if FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        if not FONT_PATH.is_file():
            raise FileNotFoundError(f"Bundled Korean font missing: {FONT_PATH}")
        pdfmetrics.registerFont(TTFont(FONT_NAME, str(FONT_PATH)))


def _label(path: list[str | int], domain: str) -> str:
    labels = dict(FALLBACK_LABELS)
    try:
        labels.update(import_module("scenarios." + domain.lower()).LABELS)
    except (ImportError, AttributeError, TypeError):
        pass
    parts = [f"{x + 1}번" if isinstance(x, int) else str(labels.get(x, x)) for x in path]
    if domain == "D_extract" and path and path[0] == "vendors" and path[-1] == "name":
        parts[-1] = "업체명"
    return " / ".join(parts) or "자료"


def _flatten(value: Any, path: list[str | int], containers: list, records: list) -> None:
    if isinstance(value, dict):
        if not all(isinstance(k, str) for k in value):
            raise TypeError("Scenario object keys must be strings")
        containers.append({"path": path, "kind": "dict"})
        if not value:
            records.append({"path": path, "kind": "dict", "value": value})
        for key, child in value.items():
            _flatten(child, path + [key], containers, records)
    elif isinstance(value, list):
        containers.append({"path": path, "kind": "list"})
        if not value:
            records.append({"path": path, "kind": "list", "value": value})
        for index, child in enumerate(value):
            _flatten(child, path + [index], containers, records)
    else:
        kind = "null" if value is None else type(value).__name__
        if kind not in ("null", "str", "int", "float", "bool"):
            raise TypeError(f"Unsupported scenario scalar: {kind}")
        if kind == "float" and not math.isfinite(value):
            raise ValueError("Scenario float must be finite")
        records.append({"path": path, "kind": kind, "value": value})


def _display(record: dict) -> str:
    kind, value = record["kind"], record["value"]
    if kind in EMPTY_LABEL:
        return EMPTY_LABEL[kind]
    if kind == "bool":
        return "예" if value else "아니요"
    if kind == "int":
        return f"{value:,}"
    if kind == "str":
        return "\u201c" + value + "\u201d"
    return repr(value)


def _wrapped(text: str, width: float, size: float) -> list[str]:
    """Split at visible glyph boundaries, preserving every original character."""
    lines = []
    for original in text.split("\n"):
        current = ""
        for char in original:
            if current and pdfmetrics.stringWidth(current + char, FONT_NAME, size) > width:
                lines.append(current)
                current = ""
            current += char
        lines.append(current)
    return lines


def _pdf(path: Path, records: list[dict], title: str, variant: int, note: str) -> list[dict]:
    _font()
    design = TEMPLATES[variant]
    accent = colors.HexColor("#" + design["accent"])
    header_fill = colors.HexColor("#" + design["header"])
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=1, pageCompression=1)
    pdf.setTitle(title)
    pdf.setAuthor("KrRubberStamp synthetic office")
    pdf.setSubject("")
    selectors = []
    page_number = 0
    margin = 38 if variant != 1 else 46
    label_width = 272 if variant != 3 else 292
    x_value = margin + label_width
    value_width = PAGE_WIDTH - margin - x_value
    size = 9.4

    def start_page() -> float:
        pdf.setFillColor(accent)
        if variant == 1:
            pdf.rect(22, 35, 7, PAGE_HEIGHT - 70, fill=1, stroke=0)
        elif variant == 4:
            pdf.roundRect(
                margin, PAGE_HEIGHT - 79, PAGE_WIDTH - 2 * margin, 42, 6, fill=1, stroke=0
            )
        if variant == 4:
            pdf.setFillColor(colors.white)
        pdf.setFont(FONT_NAME, 17)
        pdf.drawString(margin + (12 if variant == 4 else 0), PAGE_HEIGHT - 65, title)
        pdf.setFillColor(colors.HexColor("#525B66"))
        pdf.setFont(FONT_NAME, 8)
        pdf.drawString(
            margin, PAGE_HEIGHT - 99, "사내 업무용 자료 | 모든 이름과 식별번호는 합성 정보입니다"
        )
        pdf.drawString(margin, PAGE_HEIGHT - 113, note)
        pdf.setStrokeColor(accent)
        pdf.setLineWidth(2 if variant in (0, 2) else 0.7)
        pdf.line(margin, PAGE_HEIGHT - 125, PAGE_WIDTH - margin, PAGE_HEIGHT - 125)
        if variant == 2:
            pdf.line(margin, PAGE_HEIGHT - 129, PAGE_WIDTH - margin, PAGE_HEIGHT - 129)
        pdf.setFillColor(header_fill)
        pdf.rect(margin, PAGE_HEIGHT - 157, PAGE_WIDTH - 2 * margin, 24, fill=1, stroke=0)
        pdf.setFillColor(accent)
        pdf.setFont(FONT_NAME, 9)
        pdf.drawString(margin + 7, PAGE_HEIGHT - 150, "자료 항목")
        pdf.drawString(x_value + 7, PAGE_HEIGHT - 150, "원자료 기재 내용")
        pdf.setFillColor(colors.HexColor("#6C7580"))
        pdf.setFont(FONT_NAME, 7)
        pdf.drawString(margin, 22, f"KrRubberStamp | {design['name']} | {page_number + 1}쪽")
        return PAGE_HEIGHT - 157

    y = start_page()
    for record in records:
        label_lines = _wrapped(record["label"], label_width - 14, size)
        value_lines = _wrapped(_display(record), value_width - 14, size)
        row_height = max(27, 14 + 13 * max(len(label_lines), len(value_lines)))
        if y - row_height < 43:
            pdf.showPage()
            page_number += 1
            y = start_page()
        if row_height > PAGE_HEIGHT - 210:
            raise ValueError("One document field is too large to fit on a PDF page")
        if variant in (0, 1, 4) and len(selectors) % 2:
            pdf.setFillColor(colors.HexColor("#F8F9FA"))
            pdf.rect(margin, y - row_height, PAGE_WIDTH - 2 * margin, row_height, stroke=0, fill=1)
        pdf.setStrokeColor(colors.HexColor("#CCD2D8"))
        pdf.setLineWidth(0.4)
        pdf.line(margin, y - row_height, PAGE_WIDTH - margin, y - row_height)
        if variant in (0, 2, 3):
            pdf.line(x_value, y, x_value, y - row_height)
        pdf.setFillColor(colors.HexColor("#202A34"))
        pdf.setFont(FONT_NAME, size)
        for i, line in enumerate(label_lines):
            pdf.drawString(margin + 7, y - 16 - i * 13, line)
        for i, line in enumerate(value_lines):
            pdf.drawString(x_value + 7, y - 16 - i * 13, line)
        selectors.append(
            {
                "page": page_number,
                "bbox": [
                    round(x_value + 3, 3),
                    round(PAGE_HEIGHT - y, 3),
                    round(PAGE_WIDTH - margin - 3, 3),
                    round(PAGE_HEIGHT - y + row_height, 3),
                ],
            }
        )
        y -= row_height
    pdf.save()
    return selectors


def _xlsx(path: Path, records: list[dict], title: str, variant: int, note: str) -> list[dict]:
    design = TEMPLATES[variant]
    wb = Workbook()
    wb.properties.creator = "KrRubberStamp"
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    ws = wb.active
    ws.title = "원자료"
    ws.merge_cells("A1:C1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Noto Sans KR", size=18, bold=True, color=design["accent"])
    ws.merge_cells("A2:C2")
    ws["A2"] = "모든 이름과 식별번호는 합성 정보입니다"
    ws.merge_cells("A3:C3")
    ws["A3"] = note
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[3].height = 30
    ws.append([])
    for col, text in enumerate(("순번", "자료 항목", "원자료 기재 내용"), start=1):
        cell = ws.cell(5, col, text)
        cell.fill = PatternFill("solid", fgColor=design["accent"])
        cell.font = Font(name="Noto Sans KR", bold=True, color="FFFFFF", size=10)
    selectors = []
    thin = Side(style="thin", color="CFD5DB")
    for index, record in enumerate(records, start=6):
        ws.cell(index, 1, index - 5)
        ws.cell(index, 2, record["label"])
        value = record["value"]
        if record["kind"] in ("null", "dict", "list"):
            value = _display(record)
        cell = ws.cell(index, 3, value)
        # User-entered text is a cell string, never an executable formula.
        if isinstance(value, str):
            cell.data_type = "s"
        if record["kind"] == "int":
            cell.number_format = "#,##0"
        for col in range(1, 4):
            current = ws.cell(index, col)
            current.font = Font(name="Noto Sans KR", size=10, color="202A34")
            current.alignment = Alignment(vertical="top", wrap_text=True)
            current.border = Border(bottom=thin, right=thin if variant in (0, 2, 3) else Side())
            if index % 2 and variant != 2:
                current.fill = PatternFill("solid", fgColor=design["header"])
        ws.row_dimensions[index].height = max(
            26, 16 * (len(record["label"]) // 44 + 1), 16 * (str(value).count("\n") + 1)
        )
        selectors.append({"sheet": "원자료", "cell": cell.coordinate})
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 65 if variant != 3 else 70
    ws.column_dimensions["C"].width = 42
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:C{5 + len(records)}"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:5"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return selectors


def _xml(node: ET.Element) -> bytes:
    return ET.tostring(node, encoding="utf-8", xml_declaration=True)


def _hwpx(path: Path, records: list[dict], title: str, variant: int, note: str) -> list[dict]:
    """HWPX/OWPML ZIP with real table cells and a native document spine."""
    section = ET.Element(f"{{{HS}}}sec")
    paragraph_id = 0
    _font()
    label_width = (29000, 27000, 30000, 31000, 28000)[variant]
    value_width = 48000 - label_width
    font_size = (10, 10.5, 9.5, 11, 10)[variant]
    rows_to_write = [("자료 항목", "원자료 기재 내용")] + [
        (r["label"], _display(r)) for r in records
    ]
    row_heights = [
        max(
            2400,
            600
            + 1300
            * max(
                len(_wrapped(label, label_width / 100 - 6, font_size)),
                len(_wrapped(value, value_width / 100 - 6, font_size)),
            ),
        )
        for label, value in rows_to_write
    ]

    def paragraph(parent: ET.Element, text: str, *, first: bool = False) -> ET.Element:
        nonlocal paragraph_id
        paragraph_id += 1
        p = ET.SubElement(
            parent,
            f"{{{HP}}}p",
            {
                "id": str(paragraph_id),
                "paraPrIDRef": "0",
                "styleIDRef": "0",
                "pageBreak": "0",
                "columnBreak": "0",
                "merged": "0",
            },
        )
        run = ET.SubElement(p, f"{{{HP}}}run", {"charPrIDRef": "0"})
        if first:
            sec_pr = ET.SubElement(
                run,
                f"{{{HP}}}secPr",
                {
                    "id": "0",
                    "textDirection": "HORIZONTAL",
                    "spaceColumns": "0",
                    "tabStop": "8000",
                    "outlineShapeIDRef": "0",
                    "memoShapeIDRef": "0",
                    "textVerticalWidthHead": "0",
                },
            )
            ET.SubElement(
                sec_pr, f"{{{HP}}}grid", {"lineGrid": "0", "charGrid": "0", "wonggojiFormat": "0"}
            )
            ET.SubElement(
                sec_pr,
                f"{{{HP}}}startNum",
                {"pageStartsOn": "BOTH", "page": "0", "pic": "0", "tbl": "0", "equation": "0"},
            )
            ET.SubElement(
                sec_pr,
                f"{{{HP}}}visibility",
                {
                    "hideFirstHeader": "0",
                    "hideFirstFooter": "0",
                    "hideFirstMasterPage": "0",
                    "border": "SHOW_ALL",
                    "fill": "SHOW_ALL",
                    "hideFirstPageNum": "0",
                    "hideFirstEmptyLine": "0",
                    "showLineNumber": "0",
                },
            )
            page_pr = ET.SubElement(
                sec_pr,
                f"{{{HP}}}pagePr",
                {
                    "landscape": "WIDELY",
                    "width": "59528",
                    "height": "84188",
                    "gutterType": "LEFT_ONLY",
                },
            )
            ET.SubElement(
                page_pr,
                f"{{{HP}}}margin",
                {
                    "header": "4252",
                    "footer": "4252",
                    "gutter": "0",
                    "left": "5669",
                    "right": "5669",
                    "top": "5669",
                    "bottom": "5669",
                },
            )
        ET.SubElement(run, f"{{{HP}}}t").text = text
        return p

    paragraph(section, title, first=True)
    paragraph(section, "모든 이름과 식별번호는 합성 정보입니다")
    paragraph(section, note)
    paragraph(section, f"작성 양식: {TEMPLATES[variant]['name']}")
    table_parent = paragraph(section, "")
    run = table_parent.find(f"{{{HP}}}run")
    table = ET.SubElement(
        run,
        f"{{{HP}}}tbl",
        {
            "id": "1",
            "zOrder": "0",
            "numberingType": "TABLE",
            "textWrap": "TOP_AND_BOTTOM",
            "textFlow": "BOTH_SIDES",
            "lock": "0",
            "dropcapstyle": "None",
            "pageBreak": "CELL",
            "repeatHeader": "1",
            "rowCnt": str(len(records) + 1),
            "colCnt": "2",
            "cellSpacing": "0",
            "borderFillIDRef": "1",
            "noAdjust": "0",
        },
    )
    ET.SubElement(
        table,
        f"{{{HP}}}sz",
        {
            "width": "48000",
            "widthRelTo": "ABSOLUTE",
            "height": str(sum(row_heights)),
            "heightRelTo": "ABSOLUTE",
            "protect": "0",
        },
    )
    ET.SubElement(
        table,
        f"{{{HP}}}pos",
        {
            "treatAsChar": "1",
            "affectLSpacing": "0",
            "flowWithText": "1",
            "allowOverlap": "0",
            "holdAnchorAndSO": "0",
            "vertRelTo": "PARA",
            "horzRelTo": "COLUMN",
            "vertAlign": "TOP",
            "horzAlign": "LEFT",
            "vertOffset": "0",
            "horzOffset": "0",
        },
    )
    ET.SubElement(
        table, f"{{{HP}}}outMargin", {"left": "0", "right": "0", "top": "0", "bottom": "0"}
    )
    ET.SubElement(
        table, f"{{{HP}}}inMargin", {"left": "300", "right": "300", "top": "300", "bottom": "300"}
    )
    selectors = []
    for row, values in enumerate(rows_to_write):
        tr = ET.SubElement(table, f"{{{HP}}}tr")
        for col, value in enumerate(values):
            tc = ET.SubElement(
                tr,
                f"{{{HP}}}tc",
                {
                    "name": "",
                    "header": "1" if row == 0 else "0",
                    "hasMargin": "1",
                    "protect": "0",
                    "editable": "0",
                    "dirty": "0",
                    "borderFillIDRef": "1",
                },
            )
            sub = ET.SubElement(
                tc,
                f"{{{HP}}}subList",
                {
                    "id": "0",
                    "textDirection": "HORIZONTAL",
                    "lineWrap": "BREAK",
                    "vertAlign": "CENTER",
                    "linkListIDRef": "0",
                    "linkListNextIDRef": "0",
                    "textWidth": "0",
                    "textHeight": "0",
                    "hasTextRef": "0",
                    "hasNumRef": "0",
                },
            )
            paragraph(sub, value)
            ET.SubElement(tc, f"{{{HP}}}cellAddr", {"colAddr": str(col), "rowAddr": str(row)})
            ET.SubElement(tc, f"{{{HP}}}cellSpan", {"colSpan": "1", "rowSpan": "1"})
            ET.SubElement(
                tc,
                f"{{{HP}}}cellSz",
                {
                    "width": str(label_width if col == 0 else value_width),
                    "height": str(row_heights[row]),
                },
            )
            ET.SubElement(
                tc,
                f"{{{HP}}}cellMargin",
                {"left": "300", "right": "300", "top": "300", "bottom": "300"},
            )
        if row:
            selectors.append(
                {"section": "Contents/section0.xml", "table": 0, "row": row, "column": 1}
            )
    head = ET.Element(f"{{{HH}}}head", {"version": "1.4", "secCnt": "1"})
    ET.SubElement(
        head,
        f"{{{HH}}}beginNum",
        {"page": "1", "footnote": "1", "endnote": "1", "pic": "1", "tbl": "1", "equation": "1"},
    )
    refs = ET.SubElement(head, f"{{{HH}}}refList")
    fonts = ET.SubElement(refs, f"{{{HH}}}fontfaces", {"itemCnt": "7"})
    for language in ("HANGUL", "LATIN", "HANJA", "JAPANESE", "OTHER", "SYMBOL", "USER"):
        face = ET.SubElement(fonts, f"{{{HH}}}fontface", {"lang": language, "fontCnt": "1"})
        font = ET.SubElement(
            face,
            f"{{{HH}}}font",
            {"id": "0", "face": "Noto Sans KR", "type": "TTF", "isEmbedded": "0"},
        )
        ET.SubElement(
            font,
            f"{{{HH}}}typeInfo",
            {
                "familyType": "FCAT_GOTHIC",
                "serifStyle": "0",
                "weight": "5",
                "proportion": "0",
                "contrast": "0",
                "strokeVariation": "0",
                "armStyle": "0",
                "letterform": "0",
                "midline": "0",
                "xHeight": "0",
            },
        )
    fills = ET.SubElement(refs, f"{{{HH}}}borderFills", {"itemCnt": "1"})
    border = ET.SubElement(
        fills,
        f"{{{HH}}}borderFill",
        {
            "id": "1",
            "threeD": "0",
            "shadow": "0",
            "centerLine": "NONE",
            "breakCellSeparateLine": "0",
        },
    )
    for side in ("leftBorder", "rightBorder", "topBorder", "bottomBorder"):
        ET.SubElement(
            border,
            f"{{{HH}}}{side}",
            {"type": "SOLID", "width": "0.12 mm", "color": "#" + TEMPLATES[variant]["accent"]},
        )
    chars = ET.SubElement(refs, f"{{{HH}}}charProperties", {"itemCnt": "1"})
    char = ET.SubElement(
        chars,
        f"{{{HH}}}charPr",
        {
            "id": "0",
            "height": str(int(font_size * 100)),
            "textColor": "#202A34",
            "shadeColor": "none",
            "useFontSpace": "0",
            "useKerning": "0",
            "symMark": "NONE",
            "borderFillIDRef": "1",
        },
    )
    attributes = {
        language: "0"
        for language in ("hangul", "latin", "hanja", "japanese", "other", "symbol", "user")
    }
    for tag, val in (
        ("fontRef", "0"),
        ("ratio", "100"),
        ("spacing", "0"),
        ("relSz", "100"),
        ("offset", "0"),
    ):
        ET.SubElement(char, f"{{{HH}}}{tag}", {k: val for k in attributes})
    ET.SubElement(
        char, f"{{{HH}}}underline", {"type": "NONE", "shape": "SOLID", "color": "#000000"}
    )
    ET.SubElement(char, f"{{{HH}}}strikeout", {"shape": "NONE", "color": "#000000"})
    ET.SubElement(char, f"{{{HH}}}outline", {"type": "NONE"})
    ET.SubElement(
        char,
        f"{{{HH}}}shadow",
        {"type": "NONE", "color": "#C0C0C0", "offsetX": "10", "offsetY": "10"},
    )
    tabs = ET.SubElement(refs, f"{{{HH}}}tabProperties", {"itemCnt": "1"})
    ET.SubElement(tabs, f"{{{HH}}}tabPr", {"id": "0", "autoTabLeft": "0", "autoTabRight": "0"})
    paragraphs = ET.SubElement(refs, f"{{{HH}}}paraProperties", {"itemCnt": "1"})
    pp = ET.SubElement(
        paragraphs,
        f"{{{HH}}}paraPr",
        {
            "id": "0",
            "tabPrIDRef": "0",
            "condense": "0",
            "fontLineHeight": "0",
            "snapToGrid": "1",
            "suppressLineNumbers": "0",
            "checked": "0",
        },
    )
    ET.SubElement(pp, f"{{{HH}}}align", {"horizontal": "LEFT", "vertical": "BASELINE"})
    ET.SubElement(pp, f"{{{HH}}}heading", {"type": "NONE", "idRef": "0", "level": "0"})
    ET.SubElement(
        pp,
        f"{{{HH}}}breakSetting",
        {
            "breakLatinWord": "KEEP_WORD",
            "breakNonLatinWord": "KEEP_WORD",
            "widowOrphan": "0",
            "keepWithNext": "0",
            "keepLines": "0",
            "pageBreakBefore": "0",
            "lineWrap": "BREAK",
        },
    )
    margin = ET.SubElement(pp, f"{{{HH}}}margin")
    for tag in ("intent", "left", "right", "prev", "next"):
        ET.SubElement(margin, f"{{{HC}}}{tag}", {"value": "0", "unit": "HWPUNIT"})
    ET.SubElement(
        pp, f"{{{HH}}}lineSpacing", {"type": "PERCENT", "value": "160", "unit": "HWPUNIT"}
    )
    ET.SubElement(
        pp,
        f"{{{HH}}}border",
        {
            "borderFillIDRef": "1",
            "offsetLeft": "0",
            "offsetRight": "0",
            "offsetTop": "0",
            "offsetBottom": "0",
            "connect": "0",
            "ignoreMargin": "0",
        },
    )
    styles = ET.SubElement(refs, f"{{{HH}}}styles", {"itemCnt": "1"})
    ET.SubElement(
        styles,
        f"{{{HH}}}style",
        {
            "id": "0",
            "type": "PARA",
            "name": "바탕글",
            "engName": "Normal",
            "paraPrIDRef": "0",
            "charPrIDRef": "0",
            "nextStyleIDRef": "0",
            "langID": "1042",
            "lockForm": "0",
        },
    )
    package = ET.Element(f"{{{OPF}}}package", {"version": "1.0", "unique-identifier": "doc-id"})
    metadata = ET.SubElement(package, f"{{{OPF}}}metadata")
    ET.SubElement(metadata, f"{{{OPF}}}title").text = title
    manifest = ET.SubElement(package, f"{{{OPF}}}manifest")
    for ident, href in (("header", "header.xml"), ("section0", "section0.xml")):
        ET.SubElement(
            manifest, f"{{{OPF}}}item", {"id": ident, "href": href, "media-type": "application/xml"}
        )
    spine = ET.SubElement(package, f"{{{OPF}}}spine")
    ET.SubElement(spine, f"{{{OPF}}}itemref", {"idref": "header", "linear": "yes"})
    ET.SubElement(spine, f"{{{OPF}}}itemref", {"idref": "section0", "linear": "yes"})
    container = b'<?xml version="1.0" encoding="UTF-8"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="Contents/content.hpf" media-type="application/hwpml-package+xml"/></rootfiles></container>'
    odf_manifest = b'<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/hwp+zip"/><manifest:file-entry manifest:full-path="Contents/header.xml" manifest:media-type="application/xml"/><manifest:file-entry manifest:full-path="Contents/section0.xml" manifest:media-type="application/xml"/><manifest:file-entry manifest:full-path="Contents/content.hpf" manifest:media-type="application/hwpml-package+xml"/></manifest:manifest>'
    version = b'<?xml version="1.0" encoding="UTF-8"?><hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" targetApplication="WORDPROCESSOR" major="5" minor="1" micro="0" buildNumber="1" os="1" xmlVersion="1.4" application="KrRubberStamp" appVersion="0.1"/>'
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data, compression in (
            ("mimetype", b"application/hwp+zip", zipfile.ZIP_STORED),
            ("version.xml", version, zipfile.ZIP_DEFLATED),
            ("META-INF/container.xml", container, zipfile.ZIP_DEFLATED),
            ("META-INF/manifest.xml", odf_manifest, zipfile.ZIP_DEFLATED),
            ("Contents/content.hpf", _xml(package), zipfile.ZIP_DEFLATED),
            ("Contents/header.xml", _xml(head), zipfile.ZIP_DEFLATED),
            ("Contents/section0.xml", _xml(section), zipfile.ZIP_DEFLATED),
        ):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = compression
            archive.writestr(info, data)
    return selectors


def _scan(pdf_path: Path, png_path: Path, seed: int) -> None:
    rng = random.Random(seed)
    with fitz.open(pdf_path) as doc:
        pixmaps = [
            page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), colorspace=fitz.csGRAY) for page in doc
        ]
    pages = [Image.frombytes("L", (p.width, p.height), p.samples) for p in pixmaps]
    image = Image.new("L", (max(p.width for p in pages), sum(p.height for p in pages)), 255)
    y = 0
    for page in pages:
        image.paste(page, (0, y))
        y += page.height
    image = image.rotate(
        rng.choice((-0.5, -0.3, 0.4, 0.6)),
        resample=Image.Resampling.BICUBIC,
        expand=True,
        fillcolor=255,
    )
    image = image.filter(ImageFilter.GaussianBlur(0.24))
    draw = ImageDraw.Draw(image)
    stamp_font = ImageFont.truetype(str(FONT_PATH), 16)
    draw.rectangle((52, 13, image.width - 52, 40), fill=247, outline=170)
    draw.text(
        (61, 17),
        "스캔 사본 / 원본 PDF와 같은 증빙이므로 중복 집계하지 마세요",
        font=stamp_font,
        fill=70,
    )
    # Noise stays sparse and stains stay in the margins so the source remains legible.
    for _ in range(image.width * image.height // 1100):
        x, y = rng.randrange(image.width), rng.randrange(image.height)
        if image.getpixel((x, y)) > 245:
            draw.point((x, y), fill=rng.randint(220, 245))
    draw.ellipse((8, image.height // 3, 46, image.height // 3 + 84), outline=226, width=4)
    image.save(png_path, optimize=True)


def _formats(domain: str, difficulty: str) -> list[str]:
    if difficulty == "easy":
        return ["pdf", "png"] if domain == "D_extract" else ["pdf", "xlsx"]
    if difficulty == "medium":
        return (
            ["pdf", "xlsx", "hwpx", "pdf", "png"]
            if domain == "D_extract"
            else ["pdf", "xlsx", "hwpx", "pdf", "xlsx"]
        )
    if difficulty == "hard":
        return ["pdf", "xlsx", "pdf", "xlsx", "hwpx", "pdf", "xlsx", "pdf", "png", "png"]
    raise ValueError(f"Unknown difficulty: {difficulty}")


def render_task(
    scenario: dict, task_dir: Path, domain: str, difficulty: str, seed: int
) -> list[dict]:
    """Write actual source files and a private value-free extraction map.

    Scans are explicitly identified duplicate sources, never the sole source of
    a fact. Reconstruction therefore does not pretend to OCR them.
    """
    task_dir = Path(task_dir)
    input_dir = task_dir / "inputs"
    input_dir.mkdir(parents=True, exist_ok=True)
    containers, records = [], []
    _flatten(scenario, [], containers, records)
    if not records:
        raise ValueError("Scenario must contain at least one source field")
    formats = _formats(domain, difficulty)
    primary_indices = [i for i, fmt in enumerate(formats) if fmt != "png"]
    groups = {i: [] for i in primary_indices}
    blocks = {}
    for record in records:
        record["label"] = _label(record["path"], domain)
        path = record["path"]
        # Keep each transaction, person and payment together in a source sheet.
        prefix = tuple(path[:2]) if len(path) > 1 and isinstance(path[1], int) else tuple(path[:1])
        blocks.setdefault(prefix, []).append(record)
    for block in blocks.values():
        candidates = primary_indices
        # XML and workbook cells preserve source line breaks and tabs natively.
        if any(
            record["kind"] == "str" and any(char in record["value"] for char in ("\n", "\t", "\r"))
            for record in block
        ):
            candidates = [index for index in primary_indices if formats[index] != "pdf"]
            if not candidates:
                raise ValueError("Multiline source text needs an XLSX or HWPX document")
        index = min(candidates, key=lambda index: (len(groups[index]), index))
        groups[index].extend(block)
    # Small scenario smoke tests still get informative duplicate source sheets.
    for index in primary_indices:
        if not groups[index]:
            groups[index] = [records[index % len(records)]]
    input_files, documents, mapped = [], [], []
    titles = DOC_TITLES.get(domain, DOC_TITLES["D_extract"])
    pdf_indices = [i for i in primary_indices if formats[i] == "pdf"]
    for i, fmt in enumerate(formats):
        variant = (seed + i) % len(TEMPLATES)
        name = f"{i + 1:02d}_{('scan' if fmt == 'png' else 'source')}.{fmt}"
        path = input_dir / name
        relpath = f"inputs/{name}"
        title = titles[i % len(titles)]
        note = f"증빙 {i + 1:02d} | 원자료의 기재값 기준으로 검토해 주세요"
        duplicate = None
        if fmt == "png":
            source_index = pdf_indices[(i - len(primary_indices)) % len(pdf_indices)]
            duplicate = documents[source_index]["path"]
            source = task_dir / duplicate
            _scan(source, path, seed * 31 + i)
            note = "같은 거래의 스캔 사본입니다. 원본 PDF와 중복 집계하지 마세요."
        else:
            group = groups[i]
            writer = {"pdf": _pdf, "xlsx": _xlsx, "hwpx": _hwpx}[fmt]
            selectors = writer(path, group, title, variant, note)
            for record, selector in zip(group, selectors, strict=True):
                mapped.append(
                    {
                        "path": record["path"],
                        "kind": record["kind"],
                        "file": relpath,
                        "selector": selector,
                    }
                )
        entry = {"path": relpath, "format": fmt}
        if duplicate:
            entry["duplicate_of"] = duplicate
        input_files.append(entry)
        documents.append({**entry, "template_variant": variant, "form": title})
    extraction_map = {
        "version": 1,
        "containers": containers,
        "records": mapped,
        "documents": documents,
    }
    (task_dir / "extraction_map.json").write_text(
        json.dumps(extraction_map, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return input_files


def _decode(value: Any, kind: str, *, formatted: bool = False) -> Any:
    if kind == "str":
        if not formatted:
            return "" if value is None else str(value)
        if not isinstance(value, str) or not (
            value.startswith("\u201c") and value.endswith("\u201d")
        ):
            raise ValueError("Quoted text source cell is incomplete")
        return value[1:-1]
    if kind == "bool":
        if isinstance(value, bool):
            return value
        if value in ("예", "아니요"):
            return value == "예"
        raise ValueError(f"Invalid boolean source cell: {value!r}")
    if kind in EMPTY_LABEL:
        if value != EMPTY_LABEL[kind]:
            raise ValueError(f"Invalid empty source cell: {value!r}")
        return {"null": None, "dict": {}, "list": []}[kind]
    if kind == "int":
        if isinstance(value, int) and not isinstance(value, bool):
            return value
        return int(str(value).replace(",", ""))
    if kind == "float":
        return float(value)
    raise ValueError(f"Unknown scalar kind: {kind}")


def _put(root: Any, path: list, value: Any) -> Any:
    if not path:
        return value
    current = root
    for token in path[:-1]:
        current = current[token]
    token = path[-1]
    if isinstance(current, list):
        while len(current) <= token:
            current.append(None)
    current[token] = value
    return root


def restore_scenario(task_dir: Path) -> dict:
    """Read source cells/text; the extraction map supplies only selectors/types."""
    task_dir = Path(task_dir)
    mapping = json.loads((task_dir / "extraction_map.json").read_text(encoding="utf-8"))
    root = {}
    for container in sorted(mapping["containers"], key=lambda item: len(item["path"])):
        root = _put(root, container["path"], {} if container["kind"] == "dict" else [])
    cache = {}
    try:
        for document in mapping["documents"]:
            path = task_dir / document["path"]
            if not path.is_file():
                raise FileNotFoundError(path)
            if document["format"] == "png":
                with Image.open(path) as image:
                    image.verify()
        for record in mapping["records"]:
            filename, selector = record["file"], record["selector"]
            path = task_dir / filename
            if filename not in cache:
                if path.suffix == ".pdf":
                    cache[filename] = fitz.open(path)
                elif path.suffix == ".xlsx":
                    cache[filename] = load_workbook(path, read_only=True, data_only=False)
                elif path.suffix == ".hwpx":
                    with zipfile.ZipFile(path) as archive:
                        cache[filename] = ET.fromstring(archive.read(selector["section"]))
                else:
                    raise ValueError(f"Unsupported source: {filename}")
            source = cache[filename]
            if path.suffix == ".pdf":
                cell = source[selector["page"]].get_text("text", clip=fitz.Rect(selector["bbox"]))
                value = cell.rstrip("\n")
                # Numeric fields wrap only at glyph boundaries; no separators are introduced.
                if record["kind"] != "str":
                    value = value.replace("\n", "")
                elif "\n" in value:
                    value = value.replace("\n", "")
            elif path.suffix == ".xlsx":
                cell = source[selector["sheet"]][selector["cell"]]
                if cell.data_type == "f":
                    raise ValueError("Scenario source must contain values, not formulas")
                value = cell.value
            else:
                table = source.findall(f".//{{{HP}}}tbl")[selector["table"]]
                row = table.findall(f"{{{HP}}}tr")[selector["row"]]
                cell = row.findall(f"{{{HP}}}tc")[selector["column"]]
                value = "".join(node.text or "" for node in cell.iter(f"{{{HP}}}t"))
            root = _put(
                root,
                record["path"],
                _decode(value, record["kind"], formatted=path.suffix != ".xlsx"),
            )
    finally:
        for source in cache.values():
            if hasattr(source, "close"):
                source.close()
    return root
